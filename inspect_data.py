import openpyxl
import pandas as pd
import numpy as np

excel_path = "JAN26_AUG26_REPORT.xlsx"
wb = openpyxl.load_workbook(excel_path, read_only=True)
print("Sheet names:", wb.sheetnames)

all_dfs = []
for sheet in wb.sheetnames:
    df_sheet = pd.read_excel(excel_path, sheet_name=sheet)
    print(f"Sheet: {sheet}, Shape: {df_sheet.shape}")
    print("Columns:", list(df_sheet.columns)[:10])
    df_sheet['Source_Sheet'] = sheet
    all_dfs.append(df_sheet)

df = pd.concat(all_dfs, ignore_index=True)
print(f"\nTotal combined rows: {len(df)}")
print("Combined Columns:", list(df.columns))
print(df.head(3))
print("\nData info:")
print(df.info())
